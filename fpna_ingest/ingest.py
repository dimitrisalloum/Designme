#!/usr/bin/env python3
"""FP&A ingestion starter pipeline.

Converts PPTX, XLSX, and image files into normalized JSONL records and chunked text
ready for embedding / GPT workflows.
"""

from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass, asdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable, Iterator, Optional


SUPPORTED_EXTENSIONS = {
    ".pptx",
    ".xlsx",
    ".xlsm",
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".tiff",
    ".bmp",
}


@dataclass
class Record:
    doc_id: str
    source_file: str
    source_type: str
    location: str
    text: str
    period: Optional[str]
    entity: Optional[str]
    extracted_at_utc: str


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Ingest FP&A files into JSONL records.")
    parser.add_argument("input_dir", type=Path, help="Directory containing source files")
    parser.add_argument("--out-dir", type=Path, default=Path("./out"), help="Output directory")
    parser.add_argument("--glob", default="**/*", help="Glob pattern under input_dir")
    parser.add_argument("--max-chars", type=int, default=1800, help="Chunk size for text output")
    return parser.parse_args()


def discover_files(input_dir: Path, pattern: str) -> list[Path]:
    candidates = [p for p in input_dir.glob(pattern) if p.is_file()]
    return [p for p in candidates if p.suffix.lower() in SUPPORTED_EXTENSIONS]


def normalize_text(text: str) -> str:
    text = text.replace("\u00a0", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def infer_period(text: str) -> Optional[str]:
    patterns = [
        r"\b(20\d{2})[- ]?Q([1-4])\b",
        r"\bQ([1-4])[- ]?(20\d{2})\b",
        r"\b(20\d{2})[-/](0[1-9]|1[0-2])\b",
    ]
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if not match:
            continue
        groups = match.groups()
        if "Q" in pattern and len(groups) == 2:
            if groups[0].isdigit() and len(groups[0]) == 4:
                return f"{groups[0]}-Q{groups[1]}"
            return f"{groups[1]}-Q{groups[0]}"
        if len(groups) == 2:
            return f"{groups[0]}-{groups[1]}"
    return None


def infer_entity(path: Path) -> Optional[str]:
    # Starter heuristic: parent folder often contains BU/region name.
    if path.parent == path.parent.parent:
        return None
    candidate = path.parent.name
    return candidate if candidate not in {".", ".."} else None


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def extract_pptx(path: Path) -> Iterator[Record]:
    try:
        from pptx import Presentation  # type: ignore
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("python-pptx is required for .pptx parsing") from exc

    presentation = Presentation(str(path))
    for idx, slide in enumerate(presentation.slides, start=1):
        snippets: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                snippets.append(shape.text)
            if hasattr(shape, "table") and shape.has_table:
                for row in shape.table.rows:
                    snippets.append(" | ".join(cell.text for cell in row.cells))
        notes_text = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes_text = slide.notes_slide.notes_text_frame.text
            if notes_text:
                snippets.append(f"Notes: {notes_text}")

        merged = normalize_text("\n".join(s for s in snippets if s))
        if not merged:
            continue
        yield Record(
            doc_id=f"{path.stem}:slide:{idx}",
            source_file=str(path),
            source_type="ppt_slide",
            location=f"slide={idx}",
            text=merged,
            period=infer_period(merged),
            entity=infer_entity(path),
            extracted_at_utc=utc_now_iso(),
        )


def extract_xlsx(path: Path) -> Iterator[Record]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("openpyxl is required for .xlsx/.xlsm parsing") from exc

    workbook = load_workbook(filename=str(path), data_only=True, read_only=True)
    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        lines: list[str] = []
        for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                lines.append(f"r{row_index}: " + " | ".join(cells))
        merged = normalize_text("\n".join(lines))
        if not merged:
            continue
        yield Record(
            doc_id=f"{path.stem}:sheet:{sheet_name}",
            source_file=str(path),
            source_type="excel_sheet",
            location=f"sheet={sheet_name}",
            text=merged,
            period=infer_period(merged),
            entity=infer_entity(path),
            extracted_at_utc=utc_now_iso(),
        )


def extract_image_ocr(path: Path) -> Iterator[Record]:
    try:
        import pytesseract  # type: ignore
        from PIL import Image  # type: ignore
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("pytesseract and Pillow are required for image OCR") from exc

    image = Image.open(path)
    text = normalize_text(pytesseract.image_to_string(image))
    if not text:
        return
    yield Record(
        doc_id=f"{path.stem}:image:1",
        source_file=str(path),
        source_type="image_ocr",
        location="image=1",
        text=text,
        period=infer_period(text),
        entity=infer_entity(path),
        extracted_at_utc=utc_now_iso(),
    )


def extract_pdf(path: Path) -> Iterator[Record]:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("pypdf is required for .pdf parsing") from exc

    reader = PdfReader(str(path))
    for idx, page in enumerate(reader.pages, start=1):
        raw_text = page.extract_text() or ""
        text = normalize_text(raw_text)
        if not text:
            continue
        yield Record(
            doc_id=f"{path.stem}:pdf_page:{idx}",
            source_file=str(path),
            source_type="pdf_page",
            location=f"page={idx}",
            text=text,
            period=infer_period(text),
            entity=infer_entity(path),
            extracted_at_utc=utc_now_iso(),
        )


def extract_records(path: Path) -> Iterable[Record]:
    ext = path.suffix.lower()
    if ext == ".pptx":
        return extract_pptx(path)
    if ext in {".xlsx", ".xlsm"}:
        return extract_xlsx(path)
    if ext == ".pdf":
        return extract_pdf(path)
    if ext in {".png", ".jpg", ".jpeg", ".tiff", ".bmp"}:
        return extract_image_ocr(path)
    return []


def chunk_text(text: str, max_chars: int) -> list[str]:
    if len(text) <= max_chars:
        return [text]
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + max_chars, len(text))
        if end < len(text):
            split = text.rfind(" ", start, end)
            if split > start + 200:
                end = split
        chunks.append(text[start:end].strip())
        start = end
    return [c for c in chunks if c]


def write_jsonl(path: Path, rows: Iterable[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")


def main() -> None:
    args = parse_args()
    files = discover_files(args.input_dir, args.glob)
    all_records: list[Record] = []

    for file_path in files:
        try:
            extracted = list(extract_records(file_path))
            all_records.extend(extracted)
            print(f"[ok] {file_path} -> {len(extracted)} records")
        except Exception as exc:
            print(f"[warn] {file_path}: {exc}")

    record_rows = [asdict(record) for record in all_records]

    chunk_rows: list[dict] = []
    for record in all_records:
        for idx, chunk in enumerate(chunk_text(record.text, args.max_chars), start=1):
            chunk_rows.append(
                {
                    "chunk_id": f"{record.doc_id}:chunk:{idx}",
                    "doc_id": record.doc_id,
                    "source_file": record.source_file,
                    "source_type": record.source_type,
                    "location": record.location,
                    "chunk_index": idx,
                    "text": chunk,
                    "period": record.period,
                    "entity": record.entity,
                }
            )

    out_records = args.out_dir / "records.jsonl"
    out_chunks = args.out_dir / "chunks.jsonl"
    write_jsonl(out_records, record_rows)
    write_jsonl(out_chunks, chunk_rows)

    print(f"Wrote {len(record_rows)} records -> {out_records}")
    print(f"Wrote {len(chunk_rows)} chunks  -> {out_chunks}")


if __name__ == "__main__":
    main()
